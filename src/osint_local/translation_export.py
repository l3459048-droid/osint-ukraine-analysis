from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt


PAGE_HEADING_RE = re.compile(r"^## Page (\d+)\s*$", re.MULTILINE)


def export_translation_docx(
    settings,
    db,
    sha256: str,
    source_lang: str,
    target_lang: str = "ru",
) -> Path:
    row = db.get_translation(sha256, source_lang, target_lang)
    if not row:
        raise RuntimeError("Translation not found")

    markdown_path = Path(row["output_path"]).expanduser().resolve()
    translations_root = settings.translations_dir.resolve()
    try:
        markdown_path.relative_to(translations_root)
    except ValueError as exc:
        raise RuntimeError("Invalid translation path") from exc
    if not markdown_path.is_file():
        raise RuntimeError("Translation file is missing")

    document_row = db.get_document(sha256)
    if not document_row:
        raise RuntimeError("Document not found")

    markdown = markdown_path.read_text(encoding="utf-8", errors="replace")
    title, units = _parse_translation_markdown(markdown)

    document = Document()
    document.core_properties.title = title
    document.core_properties.subject = (
        f"OSINT Local translation {source_lang} → {target_lang}"
    )
    document.core_properties.author = "OSINT Local"
    document.styles["Normal"].font.size = Pt(11)

    document.add_heading(title, level=0)
    for label, value in (
        ("Source", document_row["source_path"]),
        ("SHA-256", sha256),
        ("Language", f"{source_lang} → {target_lang}"),
        ("Generated", row["created_at"] or ""),
        ("Engine", row["engine"] or "unknown"),
    ):
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{label}: ").bold = True
        paragraph.add_run(str(value))

    if units:
        document.add_paragraph()

    first_page = True
    for page, text in units:
        if page is not None:
            if not first_page:
                document.add_page_break()
            document.add_heading(f"Page {page}", level=1)
            first_page = False
        _append_text(document, text)

    output_path = markdown_path.with_suffix(".docx")
    tmp_path = output_path.with_name(output_path.name + ".tmp")
    document.save(tmp_path)
    tmp_path.replace(output_path)
    return output_path


def _parse_translation_markdown(
    markdown: str,
) -> tuple[str, list[tuple[int | None, str]]]:
    lines = markdown.splitlines()
    title = "Translation"
    if lines and lines[0].startswith("# "):
        title = lines[0][2:].strip() or title

    body_start = 1
    for index, line in enumerate(lines):
        if line.startswith("- Generated:"):
            body_start = index + 1
            break
    while body_start < len(lines) and not lines[body_start].strip():
        body_start += 1

    body = "\n".join(lines[body_start:]).strip()
    if not body:
        return title, []

    matches = list(PAGE_HEADING_RE.finditer(body))
    if not matches:
        return title, [(None, body)]

    units: list[tuple[int | None, str]] = []
    prefix = body[:matches[0].start()].strip()
    if prefix:
        units.append((None, prefix))
    for index, match in enumerate(matches):
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(body)
        text = body[start:end].strip()
        if text:
            units.append((int(match.group(1)), text))
    return title, units


def _append_text(document: Document, text: str) -> None:
    blocks = [
        block.strip()
        for block in re.split(r"\n\s*\n", text)
        if block.strip()
    ]
    for block in blocks:
        paragraph = document.add_paragraph()
        for index, line in enumerate(block.splitlines()):
            if index:
                paragraph.add_run().add_break(WD_BREAK.LINE)
            paragraph.add_run(line)



def export_translation_layout_pdf(
    settings,
    db,
    sha256: str,
    source_lang: str,
    target_lang: str = "ru",
) -> Path:
    """Create a translated PDF while preserving the original page geometry."""
    import fitz

    row = db.get_translation(sha256, source_lang, target_lang)
    if not row:
        raise RuntimeError("Translation not found")

    markdown_path = Path(row["output_path"]).expanduser().resolve()
    translations_root = settings.translations_dir.resolve()
    try:
        markdown_path.relative_to(translations_root)
    except ValueError as exc:
        raise RuntimeError("Invalid translation path") from exc
    if not markdown_path.is_file():
        raise RuntimeError("Translation file is missing")

    document_row = db.get_document(sha256)
    if not document_row:
        raise RuntimeError("Document not found")
    if str(document_row["extension"] or "").casefold() != ".pdf":
        raise RuntimeError("Layout PDF export is available only for PDF sources")

    layout_translation_path = markdown_path.with_suffix(".layout.json")
    if not layout_translation_path.is_file():
        raise RuntimeError(
            "Layout translation is missing. Re-run this PDF with the Quality engine first."
        )
    try:
        translated_layout = json.loads(
            layout_translation_path.read_text(encoding="utf-8")
        )
    except (OSError, json.JSONDecodeError) as exc:
        raise RuntimeError("Layout translation artifact is invalid") from exc

    if (
        translated_layout.get("document_sha256") != sha256
        or translated_layout.get("source_lang") != source_lang
        or translated_layout.get("target_lang") != target_lang
    ):
        raise RuntimeError("Layout translation artifact does not match this translation")

    source_path = (settings.input_dir / str(document_row["source_path"])).resolve()
    try:
        source_path.relative_to(settings.input_dir.resolve())
    except ValueError as exc:
        raise RuntimeError("Invalid source PDF path") from exc
    if not source_path.is_file():
        raise RuntimeError("Source PDF is missing")

    output_path = markdown_path.with_suffix(".layout.pdf")
    tmp_path = output_path.with_name(output_path.stem + ".tmp.pdf")
    font_regular = _find_unicode_font(bold=False)
    font_bold = _find_unicode_font(bold=True) or font_regular

    doc = fitz.open(source_path)
    stats = {
        "pages": 0,
        "blocks": 0,
        "translated_blocks": 0,
        "shrunk_blocks": 0,
        "overflow_blocks": 0,
        "font_regular": str(font_regular) if font_regular else "Base14-Cyrillic",
        "font_bold": str(font_bold) if font_bold else "Base14-Cyrillic",
    }
    try:
        page_map = {
            int(page.get("page") or 0): page
            for page in translated_layout.get("pages") or []
            if isinstance(page, dict)
        }
        if not page_map:
            raise RuntimeError("Layout translation artifact has no pages")

        for page_number, translated_page in sorted(page_map.items()):
            if page_number < 1 or page_number > len(doc):
                continue
            page = doc[page_number - 1]
            blocks = [
                block
                for block in (translated_page.get("blocks") or [])
                if isinstance(block, dict)
                and len(block.get("bbox") or []) >= 4
                and str(block.get("translated_text") or "").strip()
            ]
            if not blocks:
                continue

            placements = []
            for block in blocks:
                placement = _plan_layout_text(
                    page,
                    block,
                    blocks,
                    font_regular=font_regular,
                    font_bold=font_bold,
                )
                if placement is None:
                    stats["overflow_blocks"] += 1
                    continue
                placements.append((block, placement))
                stats["translated_blocks"] += 1
                if placement["fontsize"] + 0.05 < placement["source_fontsize"]:
                    stats["shrunk_blocks"] += 1

            # Remove only original text. Vector lines and images are preserved;
            # transparent redactions avoid painting white boxes over form lines.
            for block, placement in placements:
                rect = fitz.Rect(block["bbox"])
                rect = rect + (-0.6, -0.6, 0.6, 0.6)
                page.add_redact_annot(rect, fill=False, cross_out=False)
            if placements:
                page.apply_redactions(images=0, graphics=0, text=0)

            for block, placement in placements:
                _commit_layout_text(page, placement)

            stats["pages"] += 1
            stats["blocks"] += len(blocks)

        if tmp_path.exists():
            tmp_path.unlink()
        doc.save(tmp_path, garbage=3, deflate=True)
    finally:
        doc.close()

    tmp_path.replace(output_path)
    meta_path = output_path.with_suffix(output_path.suffix + ".json")
    meta_tmp = meta_path.with_suffix(meta_path.suffix + ".tmp")
    meta_tmp.write_text(
        json.dumps(
            {
                "version": 1,
                "document_sha256": sha256,
                "source_lang": source_lang,
                "target_lang": target_lang,
                "engine": translated_layout.get("engine") or row["engine"] or "",
                "source_pdf": str(source_path),
                "layout_translation": str(layout_translation_path),
                "output_path": str(output_path),
                **stats,
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    meta_tmp.replace(meta_path)
    return output_path


def _find_unicode_font(*, bold: bool) -> Path | None:
    candidates = (
        [
            Path("C:/Windows/Fonts/arialbd.ttf"),
            Path("C:/Windows/Fonts/segoeuib.ttf"),
            Path("C:/Windows/Fonts/calibrib.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
            Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"),
            Path("/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
        ]
        if bold
        else [
            Path("C:/Windows/Fonts/arial.ttf"),
            Path("C:/Windows/Fonts/segoeui.ttf"),
            Path("C:/Windows/Fonts/calibri.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
            Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
        ]
    )
    return next((path for path in candidates if path.is_file()), None)


def _block_style(block: dict) -> dict:
    spans = []
    for line in block.get("lines") or []:
        if isinstance(line, dict):
            spans.extend(
                span
                for span in (line.get("spans") or [])
                if isinstance(span, dict)
            )

    sizes = [
        float(span.get("size") or 0.0)
        for span in spans
        if float(span.get("size") or 0.0) > 0
    ]
    source_fontsize = max(4.0, max(sizes) if sizes else 10.0)
    flags = 0
    color = 0
    if spans:
        flags = int(spans[0].get("flags") or 0)
        color = int(spans[0].get("color") or 0)
    return {
        "source_fontsize": source_fontsize,
        "bold": bool(flags & 16),
        "serif": bool(flags & 4),
        "color": _pdf_color(color),
    }


def _pdf_color(value: int) -> tuple[float, float, float]:
    value = int(value or 0)
    return (
        ((value >> 16) & 255) / 255.0,
        ((value >> 8) & 255) / 255.0,
        (value & 255) / 255.0,
    )


def _horizontal_overlap_ratio(a, b) -> float:
    left = max(float(a.x0), float(b.x0))
    right = min(float(a.x1), float(b.x1))
    overlap = max(0.0, right - left)
    width = max(1.0, min(float(a.width), float(b.width)))
    return overlap / width


def _expanded_block_rect(page, block: dict, blocks: list[dict]):
    import fitz

    rect = fitz.Rect(block["bbox"])
    page_rect = page.rect
    nearest = float(page_rect.y1) - 4.0

    for other in blocks:
        if other is block or len(other.get("bbox") or []) < 4:
            continue
        other_rect = fitz.Rect(other["bbox"])
        if other_rect.y0 <= rect.y0 + 0.5:
            continue
        if _horizontal_overlap_ratio(rect, other_rect) < 0.20:
            continue
        nearest = min(nearest, float(other_rect.y0) - 1.0)

    original_height = max(2.0, float(rect.height))
    max_growth = max(8.0, original_height * 3.0)
    bottom = min(nearest, float(rect.y1) + max_growth)
    if bottom > rect.y1:
        rect.y1 = bottom

    rect.x0 = max(float(page_rect.x0), float(rect.x0) - 1.0)
    rect.x1 = min(float(page_rect.x1), float(rect.x1) + 1.0)
    rect.y0 = max(float(page_rect.y0), float(rect.y0) - 0.5)
    rect.y1 = min(float(page_rect.y1), float(rect.y1) + 0.5)
    return rect


def _infer_alignment(page, rect) -> int:
    import fitz

    page_width = max(1.0, float(page.rect.width))
    center_delta = abs(
        (float(rect.x0) + float(rect.x1)) / 2.0
        - page_width / 2.0
    )
    if (
        float(rect.width) <= page_width * 0.82
        and center_delta <= page_width * 0.07
    ):
        return fitz.TEXT_ALIGN_CENTER
    return fitz.TEXT_ALIGN_LEFT


def _plan_layout_text(
    page,
    block: dict,
    blocks: list[dict],
    *,
    font_regular: Path | None,
    font_bold: Path | None,
) -> dict | None:
    import fitz

    text = str(block.get("translated_text") or "").strip()
    if not text:
        return None

    style = _block_style(block)
    rect = _expanded_block_rect(page, block, blocks)
    source_fontsize = float(style["source_fontsize"])
    start_size = min(24.0, max(4.0, source_fontsize))
    min_size = max(3.2, min(5.0, start_size * 0.50))
    fontfile = font_bold if style["bold"] else font_regular
    fontname = "osintbold" if style["bold"] else "osintregular"
    if fontfile is None:
        fontname = "hebo" if style["bold"] else "helv"

    encoding = getattr(fitz, "TEXT_ENCODING_CYRILLIC", 2)
    align = _infer_alignment(page, rect)

    # PyMuPDF can wrap an over-wide single word by characters. Pre-shrink
    # against actual font metrics so Russian words stay intact when possible.
    try:
        measure_font = (
            fitz.Font(fontfile=str(fontfile))
            if fontfile is not None
            else fitz.Font(fontname=fontname)
        )
        words = re.findall(r"\S+", text)
        widest = max(
            (measure_font.text_length(word, fontsize=start_size) for word in words),
            default=0.0,
        )
        available_width = max(1.0, float(rect.width) - 1.0)
        if widest > available_width:
            start_size = max(
                min_size,
                start_size * (available_width / widest) * 0.97,
            )
    except Exception:
        pass

    size = start_size
    while size >= min_size - 0.01:
        shape = page.new_shape()
        kwargs = {
            "fontname": fontname,
            "fontsize": size,
            "color": style["color"],
            "align": align,
            "lineheight": 1.08,
        }
        if fontfile is not None:
            kwargs["fontfile"] = str(fontfile)
        else:
            kwargs["encoding"] = encoding

        rc = shape.insert_textbox(rect, text, **kwargs)
        if rc >= 0:
            return {
                "rect": rect,
                "text": text,
                "fontname": fontname,
                "fontfile": fontfile,
                "fontsize": size,
                "source_fontsize": source_fontsize,
                "color": style["color"],
                "align": align,
                "encoding": encoding,
            }
        size -= 0.5
    return None


def _commit_layout_text(page, placement: dict) -> None:
    shape = page.new_shape()
    kwargs = {
        "fontname": placement["fontname"],
        "fontsize": placement["fontsize"],
        "color": placement["color"],
        "align": placement["align"],
        "lineheight": 1.08,
    }
    if placement["fontfile"] is not None:
        kwargs["fontfile"] = str(placement["fontfile"])
    else:
        kwargs["encoding"] = placement["encoding"]

    rc = shape.insert_textbox(
        placement["rect"],
        placement["text"],
        **kwargs,
    )
    if rc >= 0:
        shape.commit(overlay=True)
